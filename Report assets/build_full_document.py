# build_full_document.py
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_helpers import add_page_number_field

from sections_front import (
    add_title_page,
    add_bonafide_certificate,
    add_acknowledgement,
    add_abstract,
    add_table_of_contents,
    add_list_of_tables,
    add_list_of_figures,
    add_list_of_abbreviations
)
from sections_intro_lit import add_chapter_1, add_chapter_2
from sections_design_thinking import add_chapter_3
from sections_testing_maintenance import add_chapter_4
from sections_results_conclusion import add_chapter_5_and_6
from sections_annexure import add_chapter_7

def build_report():
    print("Initializing Word Document...")
    doc = docx.Document()

    # Configure Default Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)

    # Configure Section Margins for A4
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25) # Standard binding margin per guidelines
    section.right_margin = Inches(1.0)

    # Configure Footer with Dynamic Page Numbering
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr = p_ftr.add_run()
    r_ftr.font.name = 'Times New Roman'
    r_ftr.font.size = Pt(11)
    add_page_number_field(r_ftr)

    # Build Document Sections
    print("Adding Title / Cover Page...")
    add_title_page(doc)

    print("Adding Bonafide Certificate...")
    add_bonafide_certificate(doc)

    print("Adding Acknowledgement...")
    add_acknowledgement(doc)

    print("Adding Abstract...")
    add_abstract(doc)

    print("Adding Table of Contents...")
    add_table_of_contents(doc)

    print("Adding List of Tables...")
    add_list_of_tables(doc)

    print("Adding List of Figures...")
    add_list_of_figures(doc)

    print("Adding List of Symbols & Abbreviations...")
    add_list_of_abbreviations(doc)

    print("Adding Chapter 1 (Introduction)...")
    add_chapter_1(doc)

    print("Adding Chapter 2 (Literature Survey)...")
    add_chapter_2(doc)

    print("Adding Chapter 3 (Design Thinking - Empathy, Define, Ideate, Prototype)...")
    add_chapter_3(doc)

    print("Adding Chapter 4 (Testing and Maintenance)...")
    add_chapter_4(doc)

    print("Adding Chapter 5 (Result) & Chapter 6 (Conclusion & Future Work)...")
    add_chapter_5_and_6(doc)

    print("Adding Chapter 7 (Annexure - Appendix I, II, III)...")
    add_chapter_7(doc)

    output_path = "ODD_GEN_Mini_Project_Report.docx"
    print(f"Saving final document to: {output_path}")
    doc.save(output_path)
    print("SUCCESS: ODD GEN Mini Project Report generated successfully!")

if __name__ == "__main__":
    build_report()
